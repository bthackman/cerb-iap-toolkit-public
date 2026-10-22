#!/usr/bin/env python3
"""
build.py — generates the CERB "Project Descriptions — Role of Intern" docx.

Reads a period's content.json and writes
CERB_Project_Descriptions_DRAFT_{period_end_date}.docx (e.g.
CERB_Project_Descriptions_DRAFT_2027-06-30.docx) into the export folder for that
period (exports/{end_date}/ unless --out-dir is given). The date always comes from
identification.json's period_to_year/month/day -- --identification, else the file
next to content.json, else its work/ subfolder, else pipeline/identification.json
(with a warning) -- see _period_end_date().

Also snapshots content.json alongside the rendered docx in the export folder (skipped
when it already lives there) -- content.json is the only structured copy of the
narrative and flags, so the export folder, the durable per-period archive, keeps it
next to the docx rendered from it.

Categories always print in ascending CERB category order (1 -> 17) within each
project, regardless of their order in content.json. If hours-grid JSON is supplied
(--grid) or found next to content.json or in its work/ subfolder (grid_1.json,
grid_2.json... for multi-page periods, else grid.json), each project's per-category
hours in content.json are checked against it and any mismatch is printed as a
WARNING -- the docx is still built, since content.json is hand-edited and the grid is
the arithmetic source of truth for the PDF. Every project should carry a
project_brief; a missing one is warned about.

    python build.py --content ../../exports/2027-06-30/content.json [--grid grid_1.json grid_2.json]
"""

import argparse
import json
import os
import shutil
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Twips
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---- Styling constants (see skills/cerb-submission-builder/references/docx-output-template.md) ----
COLOR_PROJECT   = RGBColor(0x1F, 0x4E, 0x5F)
COLOR_META      = RGBColor(0x59, 0x59, 0x59)
COLOR_CATEGORY  = RGBColor(0x2E, 0x5C, 0x7A)
COLOR_FLAG      = RGBColor(0xB2, 0x3A, 0x00)
COLOR_EXCLUSION = RGBColor(0xC0, 0x00, 0x00)
FONT_NAME = "Arial"

HERE = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = Path(HERE).parent.parent
PIPELINE_DIR = REPO_ROOT / "pipeline"


def _exports_root():
    """exports_dir from pipeline/settings.py (falls back to <repo>/exports)."""
    sys.path.insert(0, str(PIPELINE_DIR))
    try:
        from settings import exports_dir
        return exports_dir()
    except ImportError:
        return REPO_ROOT / "exports"


def set_paragraph_spacing(para, before_pt=0, after_pt=4):
    para.paragraph_format.space_before = Pt(before_pt)
    para.paragraph_format.space_after  = Pt(after_pt)


def add_bottom_border(para, color_hex):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_top_border(para, color_hex):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '6')
    top.set(qn('w:space'), '4')
    top.set(qn('w:color'), color_hex)
    pBdr.append(top)
    pPr.append(pBdr)


def run(para, text, bold=False, italic=False, size_pt=11, color=None):
    r = para.add_run(text)
    r.font.name  = FONT_NAME
    r.font.size  = Pt(size_pt)
    r.font.bold  = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = color
    return r


def bullet_para(doc, text, italic=False, color=None, indent_left=Inches(0.5), hanging=Inches(0.25)):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent    = indent_left
    para.paragraph_format.first_line_indent = -hanging
    set_paragraph_spacing(para, before_pt=0, after_pt=4)
    bullet_run = para.add_run("•\t")
    bullet_run.font.name = FONT_NAME
    bullet_run.font.size = Pt(11)
    if color:
        bullet_run.font.color.rgb = color
    run(para, text, italic=italic, size_pt=11, color=color)
    return para


def sub_bullet_para(doc, text, indent_left=Inches(0.9), hanging=Inches(0.25)):
    """Second-level 'including but not limited to:' item -- hollow-circle marker,
    indented further than a first-level bullet_para(). See docx-output-template.md."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = indent_left
    para.paragraph_format.first_line_indent = -hanging
    set_paragraph_spacing(para, before_pt=0, after_pt=4)
    bullet_run = para.add_run("o\t")
    bullet_run.font.name = FONT_NAME
    bullet_run.font.size = Pt(11)
    run(para, text, size_pt=11)
    return para


def _resolve_identification(content_path, explicit=None):
    """--identification if given; else identification.json next to content.json or in
    its work/ subfolder (a period's own folder); else pipeline/identification.json,
    with a warning, since that file has no period dates of its own."""
    if explicit:
        return Path(explicit)
    folder = Path(content_path).parent
    for candidate in (folder / "identification.json", folder / "work" / "identification.json"):
        if candidate.exists():
            return candidate
    root_fallback = PIPELINE_DIR / "identification.json"
    print(
        "WARNING: no identification.json next to {} -- falling back to {}.".format(
            content_path, root_fallback),
        file=sys.stderr,
    )
    return root_fallback


def _period_end_date(id_path):
    """Reads period_to_year/month/day from identification.json and returns it as
    'YYYY-MM-DD', or None if the file is missing or those fields aren't set. This is
    the single source of truth for dating output filenames -- see build()'s standing
    rule that every generated file name include the period end date, so a docx/PDF
    pulled out of its exports/{end_date}/ folder (e.g. emailed) still identifies
    itself."""
    if not id_path.exists():
        return None
    with open(id_path, encoding="utf-8") as f:
        ident = json.load(f)
    y = ident.get("period_to_year", "")
    m = ident.get("period_to_month", "")
    d = ident.get("period_to_day", "")
    if y and m and d:
        return "{}-{}-{}".format(y, m.zfill(2), d.zfill(2))
    return None


def _export_dir_from_identification(id_path, end_date):
    """exports/{end_date}/ (settings exports_dir), or the current folder when undated."""
    if end_date:
        export_dir = _exports_root() / end_date
        export_dir.mkdir(parents=True, exist_ok=True)
        return export_dir
    return Path.cwd()


def _category_sort_key(cat):
    try:
        return int(cat.get("number"))
    except (TypeError, ValueError):
        return 99  # unparsable numbers sort last rather than crashing the build


def _find_grid_files(content_path):
    """Grid JSON next to content.json: per-page grid_N.json files (multi-page period,
    in page order) if any exist, else a single grid.json, else none. Looks next to
    content.json first, then in its work/ subfolder."""
    for folder in (Path(content_path).parent, Path(content_path).parent / "work"):
        pages = []
        for p in folder.glob("grid_*.json"):
            suffix = p.stem[len("grid_"):]
            if suffix.isdigit():
                pages.append((int(suffix), p))
        if pages:
            return [p for _, p in sorted(pages)]
        single = folder / "grid.json"
        if single.exists():
            return [single]
    return []


def check_against_grid(content, grid_paths):
    """Compare content.json's per-project per-category hours with the hours grid(s).
    Grid columns restart at 1 on each page, so page N's columns are offset by the
    project count of the pages before it -- the same global numbering content.json
    uses. Returns a list of warning strings (empty = consistent)."""
    grid_hours = {}  # global project number -> {category number: hours}
    grid_names = {}
    offset = 0
    for path in grid_paths:
        with open(path, encoding="utf-8") as f:
            grid = json.load(f)
        for p in grid.get("projects", []):
            grid_names[offset + p["column"]] = p["name"]
            grid_hours.setdefault(offset + p["column"], {})
        for row in grid.get("category_rows", []):
            for col, hours in row.get("hours_by_project_column", {}).items():
                if hours:
                    grid_hours.setdefault(offset + int(col), {})[row["category_number"]] = round(hours, 2)
        offset += len(grid.get("projects", []))

    warnings = []
    for project in content.get("projects", []):
        number = project.get("number")
        if number not in grid_hours:
            warnings.append("Project {} ('{}') has no matching grid column.".format(number, project.get("name")))
            continue
        content_hours = {}
        for cat in project.get("categories", []):
            key = _category_sort_key(cat)
            content_hours[key] = round(content_hours.get(key, 0.0) + float(cat.get("hours") or 0), 2)
        for cat_num in sorted(set(content_hours) | set(grid_hours[number])):
            c = content_hours.get(cat_num, 0.0)
            g = grid_hours[number].get(cat_num, 0.0)
            if abs(c - g) > 0.005:
                warnings.append("Project {} ('{}' / grid '{}') category {}: content.json {} hrs vs grid {} hrs.".format(
                    number, project.get("name"), grid_names.get(number), cat_num, c, g))
    content_numbers = {p.get("number") for p in content.get("projects", [])}
    for number in sorted(set(grid_hours) - content_numbers):
        warnings.append("Grid project {} ('{}') is missing from content.json.".format(number, grid_names.get(number)))
    return warnings


def build(content_path, out_dir=None, grid_paths=None, identification_path=None):
    with open(content_path, encoding="utf-8") as f:
        content = json.load(f)

    for project in content.get("projects", []):
        if not project.get("project_brief"):
            print("WARNING: project {} ('{}') has no project_brief -- every project needs one.".format(
                project.get("number"), project.get("name")), file=sys.stderr)

    # Docx/PDF consistency check -- warn only, never block the build.
    if grid_paths is None:
        grid_paths = _find_grid_files(content_path)
    if grid_paths:
        mismatches = check_against_grid(content, grid_paths)
        for w in mismatches:
            print("WARNING: docx/grid hours mismatch -- {}".format(w), file=sys.stderr)
        if not mismatches:
            print("Hours check: content.json matches {}.".format(", ".join(Path(p).name for p in grid_paths)))

    doc = Document()

    # ---- Page setup: US Letter, 0.75" margins ----
    section = doc.sections[0]
    section.page_width  = Twips(12240)
    section.page_height = Twips(15840)
    section.top_margin    = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin   = Inches(0.75)
    section.right_margin  = Inches(0.75)

    # Remove default paragraph spacing from Normal style
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)
    doc.styles['Normal'].font.name = FONT_NAME
    doc.styles['Normal'].font.size = Pt(11)

    # ---- Title block ----
    t = doc.add_paragraph()
    run(t, content["title"], bold=True, size_pt=16)
    set_paragraph_spacing(t, before_pt=0, after_pt=4)

    s = doc.add_paragraph()
    run(s, content["subtitle"], italic=True, size_pt=10, color=COLOR_META)
    set_paragraph_spacing(s, before_pt=0, after_pt=4)

    n = doc.add_paragraph()
    run(n, content["note"], italic=True, size_pt=10, color=COLOR_FLAG)
    set_paragraph_spacing(n, before_pt=0, after_pt=16)

    # ---- Projects ----
    for i, project in enumerate(content["projects"]):

        # Page break before every project except the first
        if i > 0:
            pb = doc.add_paragraph()
            pb.add_run().add_break(WD_BREAK.PAGE)
            set_paragraph_spacing(pb, 0, 0)

        # Project heading with bottom border
        heading_text = "{}.  {}".format(project["number"], project["name"])
        h = doc.add_paragraph()
        run(h, heading_text, bold=True, size_pt=14, color=COLOR_PROJECT)
        add_bottom_border(h, "1F4E5F")
        set_paragraph_spacing(h, before_pt=12, after_pt=2)

        # Meta line (only if fields have values)
        meta_parts = [v for v in [project.get("project_type"), project.get("occupancy")] if v]
        if meta_parts:
            m = doc.add_paragraph()
            run(m, "  |  ".join(meta_parts), italic=True, size_pt=10, color=COLOR_META)
            set_paragraph_spacing(m, before_pt=0, after_pt=8)
        else:
            spacer = doc.add_paragraph()
            set_paragraph_spacing(spacer, before_pt=0, after_pt=4)

        # Project Brief / Architectural Experience Description labels -- every project
        # should have a brief (see docx-output-template.md); rendered only when present.
        if project.get("project_brief"):
            bl = doc.add_paragraph()
            run(bl, "Project Brief:", bold=True, size_pt=10, color=COLOR_META)
            set_paragraph_spacing(bl, before_pt=0, after_pt=2)

            bp = doc.add_paragraph()
            run(bp, project["project_brief"], size_pt=11)
            set_paragraph_spacing(bp, before_pt=0, after_pt=8)

            al = doc.add_paragraph()
            run(al, "Architectural Experience Description:", bold=True, size_pt=10, color=COLOR_META)
            set_paragraph_spacing(al, before_pt=0, after_pt=8)

        # Categories -- always ascending by CERB category number (1 -> 17), whatever
        # order content.json lists them in.
        for cat in sorted(project["categories"], key=_category_sort_key):
            cat_text = "{}.  {}  |  {} Hours".format(
                cat["number"], cat["name"], cat["hours"]
            )
            ch = doc.add_paragraph()
            run(ch, cat_text, bold=True, size_pt=11, color=COLOR_CATEGORY)
            set_paragraph_spacing(ch, before_pt=8, after_pt=4)

            for bullet in cat.get("bullets") or []:
                # A bullet is either a plain string, or {"text": ..., "sub_items": [...]}
                # for the "...including but not limited to:" nested-list convention.
                if isinstance(bullet, dict):
                    bullet_para(doc, bullet["text"])
                    for sub_item in bullet.get("sub_items") or []:
                        sub_bullet_para(doc, sub_item)
                else:
                    bullet_para(doc, bullet)

            for flag in cat.get("flags") or []:
                fp = doc.add_paragraph()
                fp.paragraph_format.left_indent = Inches(0.75)
                set_paragraph_spacing(fp, before_pt=2, after_pt=6)
                run(fp, "FLAGGED — " + flag, italic=True, size_pt=10, color=COLOR_FLAG)

    # ---- Exclusions section ----
    exclusions = content.get("exclusions") or []
    if exclusions:
        sep = doc.add_paragraph()
        add_top_border(sep, "C00000")
        set_paragraph_spacing(sep, before_pt=18, after_pt=4)

        eh = doc.add_paragraph()
        run(eh, "Excluded from this draft", bold=True, size_pt=12, color=COLOR_EXCLUSION)
        set_paragraph_spacing(eh, before_pt=0, after_pt=6)

        for exc in exclusions:
            ep = doc.add_paragraph()
            run(ep, exc, italic=True, size_pt=10, color=COLOR_EXCLUSION)
            set_paragraph_spacing(ep, before_pt=0, after_pt=5)

    # ---- Flags summary ----
    if content.get("flags_summary"):
        fs = doc.add_paragraph()
        add_top_border(fs, "AAAAAA")
        set_paragraph_spacing(fs, before_pt=18, after_pt=0)
        run(fs, content["flags_summary"], italic=True, size_pt=10, color=COLOR_META)

    # Standing rule: every generated docx is named with its period's end date, so the
    # file still identifies itself once it's out of its exports/{end_date}/ folder
    # (e.g. emailed) -- see _period_end_date()'s docstring. This applies regardless of
    # whether --out-dir was passed; only the directory is affected by that flag, not
    # whether the filename gets dated.
    id_path = _resolve_identification(content_path, identification_path)
    end_date = _period_end_date(id_path)

    if out_dir is None:
        export_dir = _export_dir_from_identification(id_path, end_date)
    else:
        export_dir = Path(out_dir)
        export_dir.mkdir(parents=True, exist_ok=True)

    if end_date:
        filename = "CERB_Project_Descriptions_DRAFT_{}.docx".format(end_date)
    else:
        print(
            "WARNING: no period_to_year/month/day in {} -- output filename won't be "
            "dated.".format(id_path),
            file=sys.stderr,
        )
        filename = "CERB_Project_Descriptions_DRAFT.docx"

    out_path = export_dir / filename
    doc.save(str(out_path))
    print("Written: {}".format(out_path))

    # Snapshot the narrative source into the export folder -- see module docstring.
    snapshot_path = export_dir / "content.json"
    if snapshot_path.resolve() != Path(content_path).resolve():
        shutil.copyfile(content_path, snapshot_path)
        print("Snapshot: {}".format(snapshot_path))


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument(
        "--content", required=True,
        help="Path to this period's content.json (usually exports/{end_date}/content.json)",
    )
    parser.add_argument(
        "--identification", default=None,
        help="Path to this period's identification.json (default: next to content.json, then its work/ subfolder)",
    )
    parser.add_argument(
        "--out-dir", default=None,
        help="Directory to write the docx into (default: derived from identification.json)",
    )
    parser.add_argument(
        "--grid", nargs="+", default=None,
        help="Hours-grid JSON(s) to check content.json's hours against, in page order "
             "(default: grid_N.json / grid.json next to content.json, if present)",
    )
    args = parser.parse_args()
    build(args.content, args.out_dir, args.grid, args.identification)


if __name__ == "__main__":
    main()
